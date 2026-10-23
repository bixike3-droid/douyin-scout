"""Offline verification of the actual distributed runtime, not a simulated agent."""
import asyncio
from dataclasses import replace
import json
import os
from pathlib import Path
import subprocess
import sys
import traceback


def check(paths, output):
    from ..app import AppContext
    from ..config import load_config
    from ..codex.runtime import server, executable
    from ..store.conversation import Conversations
    from ..store.attachments import Attachments
    from ..attachments import import_file, ffmpeg_exe, _video
    from ..tools import workspace
    from PIL import Image

    report = {"frozen": bool(getattr(sys, "frozen", False)), "checks": []}
    try:
        config = load_config(paths)
        config = replace(config, llm=replace(config.llm, provider="deepseek", api_key="offline-release-check", model="deepseek-flash"))
        with AppContext.create(config) as ctx:
            ctx.tool_conversation_id = Conversations(ctx.db).create("发行包运行环境检查")
            ctx.chat_run_id = None
            ctx.tool_emit = lambda event: None
            image = paths.state_dir / "release-source.png"
            Image.new("RGB", (64, 48), "#2976d3").save(image)
            item = import_file(paths, image)
            Attachments(ctx.db).add(item, ctx.tool_conversation_id)
            blocked = paths.root / "outside-workspace.txt"
            code = """
import csv, json, os, sqlite3, statistics, zipfile
from pathlib import Path
from PIL import Image
assert 'DEEPSEEK_API_KEY' not in os.environ
data=json.loads(Path(os.environ['SCOUT_INPUT_MANIFEST']).read_text(encoding='utf-8'))
with Image.open(data['attachments'][0]['path']) as image:
    image.resize((128,96)).save('result.png')
with open('结果.csv','w',encoding='utf-8',newline='') as handle:
    csv.writer(handle).writerows([['文件','宽度'],['result.png',128]])
from openpyxl import Workbook
book=Workbook()
book.active.append(['文件','宽度'])
book.active.append(['result.png',128])
book.save('结果.xlsx')
from docx import Document
document=Document()
document.add_paragraph('源镜独立脚本环境检查')
document.save('结果.docx')
from reportlab.pdfgen.canvas import Canvas
canvas=Canvas('result.pdf')
canvas.drawString(60,760,'Yuanjing standalone workspace check')
canvas.save()
try:
    Path(BLOCKED_PATH).write_text('must fail',encoding='utf-8')
except (PermissionError,OSError):
    print('sandbox-denied-outside-write')
else:
    raise AssertionError('outside write unexpectedly succeeded')
print('独立脚本运行成功')
""".replace("BLOCKED_PATH", repr(str(blocked)))

            async def run():
                async with server(paths, config.llm) as rpc:
                    ctx.codex_rpc = rpc
                    result = await workspace._python(ctx, {"code": code, "description": "运行内置 Python 与真实 Windows 沙箱"})
                    assert result["ok"], result
                    assert "独立脚本运行成功" in result["stdout"], result
                    assert not blocked.exists()
                    report["checks"].append({"name": "bundled_python_and_mxc", "passed": True})
                    media = await workspace._ffmpeg(ctx, {"arguments": ["-f", "lavfi", "-i", "color=c=blue:s=128x96:d=1", "-y", "video.mp4"], "description": "内置视频处理"})
                    assert media["ok"], media
                    frames = paths.state_dir / "release-video-frames"
                    frames.mkdir(exist_ok=True)
                    note, visuals, _ = _video(workspace.root(ctx) / "video.mp4", frames)
                    assert "1.00" in note and len(visuals) == 5
                    report["checks"].append({"name": "bundled_ffmpeg", "passed": True})
                    assert len(workspace.scan_artifacts(ctx)) >= 6
                    from ..attachments import extract
                    for name in ("结果.xlsx", "结果.docx", "result.pdf"):
                        imported = import_file(paths, workspace.root(ctx) / name)
                        parsed = extract(paths, imported)
                        assert parsed.get("status") != "failed" and parsed.get("text"), (name, parsed)
                    report["checks"].append({"name": "office_and_pdf_script_outputs", "passed": True})
                    report["checks"].append({"name": "durable_artifacts", "passed": True})
            asyncio.run(run())
            assert ctx.db.query_one("PRAGMA quick_check")[0] == "ok"
            report["checks"].append({"name": "database_migrations", "passed": True})
        from .api import create_app
        from .server import Server
        import urllib.request
        app = create_app(paths)
        carrier = Server(app)
        carrier.start()
        try:
            origin = f"http://127.0.0.1:{carrier.port}"
            with urllib.request.urlopen(origin) as response:
                assert b'assets/' in response.read()
            request = urllib.request.Request(origin + "/api/bootstrap", headers={"Authorization": "Bearer " + app.state.gateway.token})
            with urllib.request.urlopen(request) as response:
                assert json.load(response)["settings"]["model"] == "deepseek-flash"
            report["checks"].append({"name": "packaged_frontend_and_private_api", "passed": True})
        finally:
            carrier.close()
        report["verified"] = True
    except Exception:
        report["verified"] = False
        report["error"] = traceback.format_exc()
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return 0 if report["verified"] else 1
